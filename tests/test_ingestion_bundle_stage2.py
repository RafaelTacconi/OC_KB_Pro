"""
test_ingestion_bundle_stage2.py — SPEC §22.4/§22.5/§22.6.

Criteria: A69 (PDF page labelling), A70 (DOCX titles / XLSX sheet_name, page
rule PDF-only), A71 (chunks.page migration), A72 (API page — covered by the
lookup used in service/engine.py), A73 (PDF chip shows page), A74 (furniture by
repetition), A75 (furniture never empties a page), A77 (stale citation), A82
(re-ingestion script shape).

All inputs are purpose-built under pytest's tmp_path. No synthetic or real
corpus document is used, and no real corpus content appears here (AGENTS.md
rule 8).
"""

from __future__ import annotations

import importlib
import sqlite3
from pathlib import Path

from docx import Document

import db as db_module
from ingestion.chunking import chunk_text
from ingestion.parsers import parse_docx, parse_pdf, parse_xlsx
from ui.cards import STALE_CITATION_NOTE, STALE_FALLBACK_LABEL, chips_html

REPO_ROOT = Path(__file__).resolve().parents[1]

_LEGACY_CHUNKS_SQL = """
    CREATE TABLE chunks (
        chunk_id TEXT PRIMARY KEY, source_id TEXT, workspace_id TEXT,
        section_title TEXT, text TEXT NOT NULL, embedding_text TEXT,
        embedding BLOB, token_count INTEGER, created_at TEXT NOT NULL
    );
"""


# --- helpers ---------------------------------------------------------------

def _pdf_escape(s: str) -> str:
    return s.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def _make_text_pdf(path: Path, pages: list[list[str]]) -> None:
    """Build a minimal valid multi-page PDF with real text, no dependencies."""
    path = Path(path)
    objs: dict[int, bytes] = {}
    font_num = 3
    page_nums: list[int] = []
    content_nums: list[int] = []
    n = 4
    for _ in pages:
        page_nums.append(n)
        content_nums.append(n + 1)
        n += 2
    objs[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    kids = " ".join(f"{p} 0 R" for p in page_nums)
    objs[2] = f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode()
    objs[font_num] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    for i, lines in enumerate(pages):
        pnum, cnum = page_nums[i], content_nums[i]
        objs[pnum] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_num} 0 R >> >> /Contents {cnum} 0 R >>"
        ).encode()
        y = 720
        stream = []
        for line in lines:
            stream += ["BT", "/F1 12 Tf", f"72 {y} Td", f"({_pdf_escape(line)}) Tj", "ET"]
            y -= 40
        s = "\n".join(stream).encode()
        objs[cnum] = b"<< /Length " + str(len(s)).encode() + b" >>\nstream\n" + s + b"\nendstream"
    out = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for num in sorted(objs):
        offsets[num] = len(out)
        out += f"{num} 0 obj\n".encode() + objs[num] + b"\nendobj\n"
    xref_pos = len(out)
    maxnum = max(objs)
    out += f"xref\n0 {maxnum + 1}\n".encode() + b"0000000000 65535 f \n"
    for num in range(1, maxnum + 1):
        off = offsets.get(num)
        out += (f"{off:010d} 00000 n \n".encode() if off is not None
                else b"0000000000 65535 f \n")
    out += f"trailer\n<< /Size {maxnum + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))


def _columns(db_path: Path, table: str) -> set[str]:
    conn = sqlite3.connect(db_path)
    try:
        return {r[1] for r in conn.execute(f"PRAGMA table_info({table})")}
    finally:
        conn.close()


# --- A71: migration --------------------------------------------------------

def test_migration_adds_page_and_is_idempotent(tmp_path):
    legacy = tmp_path / "legacy.db"
    conn = sqlite3.connect(legacy)
    conn.executescript(_LEGACY_CHUNKS_SQL)
    conn.commit()
    conn.close()
    assert "page" not in _columns(legacy, "chunks")

    db_module.migrate_db(legacy)
    after_first = _columns(legacy, "chunks")
    assert "page" in after_first

    db_module.migrate_db(legacy)  # idempotent: second run changes nothing
    assert _columns(legacy, "chunks") == after_first


def test_fresh_and_migrated_databases_have_the_same_chunks_columns(tmp_path):
    fresh = tmp_path / "fresh.db"
    db_module.init_db(fresh)

    legacy = tmp_path / "legacy2.db"
    conn = sqlite3.connect(legacy)
    conn.executescript(_LEGACY_CHUNKS_SQL)
    conn.commit()
    conn.close()
    db_module.migrate_db(legacy)

    assert _columns(fresh, "chunks") == _columns(legacy, "chunks")


# --- A69 / A71: PDF page capture -------------------------------------------

def test_pdf_chunks_carry_their_page(tmp_path):
    pdf = tmp_path / "two_pages.pdf"
    _make_text_pdf(pdf, [["Alpha body text"], ["Beta body text"]])
    sections = parse_pdf(str(pdf))
    assert [s.page for s in sections] == [1, 2]
    assert all(s.section_title is None for s in sections)  # no inferred PDF title
    page_two = next(s for s in sections if s.page == 2)
    chunks = chunk_text(page_two.text, section_title=page_two.section_title, page=page_two.page)
    assert chunks
    assert all(c["page"] == 2 for c in chunks)


# --- A70: format split ------------------------------------------------------

def test_docx_keeps_section_titles_and_no_page(tmp_path):
    doc = Document()
    doc.add_heading("Section One", level=1)
    doc.add_paragraph("First body paragraph.")
    doc.add_heading("Section Two", level=1)
    doc.add_paragraph("Second body paragraph.")
    path = tmp_path / "doc.docx"
    doc.save(str(path))

    sections = parse_docx(str(path))
    titles = [s.section_title for s in sections]
    assert any(t and "Section One" in t for t in titles)
    assert all(s.page is None for s in sections)
    for s in sections:
        for chunk in chunk_text(s.text, section_title=s.section_title, page=s.page):
            assert chunk["page"] is None


def test_xlsx_keeps_sheet_name_and_no_page(tmp_path):
    import pandas as pd

    path = tmp_path / "book.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame([["A", "B"], ["1", "2"]]).to_excel(
            writer, sheet_name="MySheet", index=False, header=False
        )
    sections = parse_xlsx(str(path))
    assert sections[0].section_title == "MySheet"
    assert sections[0].page is None


# --- A74 / A75: page furniture ---------------------------------------------

def test_repeated_footer_is_removed(tmp_path):
    pdf = tmp_path / "footer.pdf"
    _make_text_pdf(pdf, [
        ["Alpha body", "Footer CONFIDENTIAL Page 1"],
        ["Beta body", "Footer CONFIDENTIAL Page 2"],
    ])
    text = " ".join(s.text for s in parse_pdf(str(pdf)))
    assert "Alpha body" in text and "Beta body" in text
    assert "Footer CONFIDENTIAL" not in text  # digits normalised; removed


def test_line_on_only_some_pages_is_kept(tmp_path):
    pdf = tmp_path / "partial.pdf"
    _make_text_pdf(pdf, [
        ["Alpha body", "Only on page one note"],
        ["Beta body", "A different note"],
    ])
    text = " ".join(s.text for s in parse_pdf(str(pdf)))
    assert "Only on page one note" in text


def test_single_page_document_removes_nothing(tmp_path):
    pdf = tmp_path / "single.pdf"
    _make_text_pdf(pdf, [["Alpha body", "Footer Page 1"]])
    text = " ".join(s.text for s in parse_pdf(str(pdf)))
    assert "Footer Page 1" in text


def test_furniture_retained_when_removal_would_empty_a_page(tmp_path):
    pdf = tmp_path / "emptyrisk.pdf"
    _make_text_pdf(pdf, [["REPEAT ONLY"], ["Body text", "REPEAT ONLY"]])
    sections = parse_pdf(str(pdf))
    page_one = next(s for s in sections if s.page == 1)
    assert "REPEAT ONLY" in page_one.text  # retained; the page is not lost


# --- A73: chip page label ---------------------------------------------------

def test_pdf_chip_shows_page_number():
    html = chips_html([{"display_name": "Doc.pdf", "section_title": None, "page": 3}])
    assert "Page 3" in html


def test_docx_chip_shows_section_not_page():
    html = chips_html([{"display_name": "Doc.docx", "section_title": "1. Scope", "page": None}])
    assert "1. Scope" in html
    assert "Page" not in html


# --- A77: stale citations ---------------------------------------------------

def test_stale_citation_renders_marked_and_not_clickable():
    html = chips_html(
        [{"display_name": "Old.pdf", "section_title": "2. Timeline", "page": None}],
        heading="Retrieved from",
        stale=True,
    )
    assert STALE_CITATION_NOTE in html
    assert "Old.pdf" in html          # still rendered, not silently vanished
    assert "<a " not in html          # not clickable


def test_stale_citation_fallback_label_for_unidentifiable_source():
    html = chips_html([{}], stale=True)
    assert STALE_FALLBACK_LABEL in html


def test_stored_citations_stale_helper(monkeypatch, tmp_path):
    monkeypatch.chdir(tmp_path)
    db_module.init_db()
    db_module.migrate_db()
    from ui.chat_view import _stored_citations_stale

    assert _stored_citations_stale('["does-not-exist"]') is True
    assert _stored_citations_stale(None) is False
    assert _stored_citations_stale("[]") is False
    assert _stored_citations_stale("not json") is False


# --- A72: page lookup used by the API --------------------------------------

def test_chunk_pages_lookup(monkeypatch, tmp_path):
    monkeypatch.chdir(tmp_path)
    db_module.init_db()
    db_module.migrate_db()
    conn = sqlite3.connect(tmp_path / "data" / "workspace_app.db")
    try:
        conn.execute(
            "INSERT INTO chunks (chunk_id, source_id, workspace_id, section_title, page, "
            "text, created_at) VALUES (?,?,?,?,?,?,?)",
            ("c1", "s1", "w1", None, 4, "text", "2026-09-16"),
        )
        conn.commit()
    finally:
        conn.close()
    pages = db_module.chunk_pages(["c1", "missing"])
    assert pages.get("c1") == 4
    assert "missing" not in pages


# --- A82: re-ingestion script shape ----------------------------------------

def test_reingest_script_is_a_tool_not_a_test():
    script = REPO_ROOT / "scripts" / "reingest_all.py"
    assert script.exists()
    assert script.parent.name == "scripts"      # not under tests/
    assert not script.name.startswith("test_")  # not collected by pytest
    text = script.read_text(encoding="utf-8")
    assert 'if __name__ == "__main__":' in text  # no work on import

    module = importlib.import_module("scripts.reingest_all")
    assert callable(module.main)
    assert module.WORKSPACES == [
        "aml-workspace",
        "62ee0cd0d2324cd5ab5e20b70a8b27bc",
        "aa6dc06b19dd47e1a24d4855b951d24f",
        "524bf0ec5c5f4290a80438fd801455ce",
        "d93d5cd12269400e9298c535062afc65",
    ]
